"""Tests of the `ler` reader, without Agent Zero. Run: python3 infra/vm/plugins/ler_arquivo/tests/suite.py"""

import asyncio
import importlib.util
import sys
import tempfile
import traceback
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]


class Falha(AssertionError):
    pass


def ok(cond, msg):
    if not cond:
        raise Falha(msg)


def _leitor():
    spec = importlib.util.spec_from_file_location("ler_arquivo_teste", RAIZ / "helpers" / "ler.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


class Suite:
    def __init__(self):
        self.resultados = []
        self.tmp = Path(tempfile.mkdtemp(prefix="ler_teste_"))

    def t_planilha_inteira_em_paginas_e_aba_escondida_fora(self):
        import openpyxl

        m = _leitor()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Itens"
        ws.append(["NCM", "Descrição", "Qtde"])
        for i in range(1, 451):
            ws.append([f"6809.11.{i:04d}", f"Placa {i}", i])
        ws.append(["Total", None, "=SUM(C2:C451)"])  # fórmula sem valor calculado (openpyxl não calcula)
        oculta = wb.create_sheet("Rascunho")
        oculta.sheet_state = "hidden"
        oculta.append(["segredo"])
        arq = self.tmp / "ncm.xlsx"
        wb.save(arq)
        r = m.ler(str(arq))
        ok("Itens" in r and "452 linhas" in r, f"conta todas as linhas: {r[:300]}")
        ok("L2: 6809.11.0001 | Placa 1 | 1" in r, "linha com valores, separada por |")
        ok("escondidas" in r and "segredo" not in r, "aba escondida fica de fora e é citada")
        ok("`inicio` = 201" in r, "diz como continuar")
        r2 = m.ler(str(arq), inicio=440, limite=50)
        ok("L452: Total" in r2 and "fórmula sem valor calculado" in r2, f"fim da planilha e aviso da fórmula: {r2[-400:]}")
        ok("segredo" in m.ler(str(arq), aba="Rascunho"), "aba escondida lida quando pedida")
        ok("não existe" in m.ler(str(arq), aba="Nada"), "aba que não existe")

    def t_valor_calculado_e_nao_a_formula(self):
        import openpyxl
        import zipfile

        m = _leitor()
        wb = openpyxl.Workbook()
        wb.active.append([2, 3, "=A1+B1"])
        arq = self.tmp / "conta.xlsx"
        wb.save(arq)
        # simula o Excel: grava o valor calculado (<v>5</v>) junto da fórmula
        with zipfile.ZipFile(arq) as z:
            partes = {n: z.read(n) for n in z.namelist()}
        nome = "xl/worksheets/sheet1.xml"
        partes[nome] = partes[nome].replace(b"<f>A1+B1</f><v></v>", b"<f>A1+B1</f><v>5</v>").replace(
            b"<f>A1+B1</f><v />", b"<f>A1+B1</f><v>5</v>")
        with zipfile.ZipFile(arq, "w") as z:
            for n, dados in partes.items():
                z.writestr(n, dados)
        r = m.ler(str(arq))
        ok("L1: 2 | 3 | 5" in r and "=A1+B1" not in r, f"vem o valor, não a fórmula: {r}")

    def t_csv_docx_pdf_texto(self):
        import docx
        import pypdf

        m = _leitor()
        csvf = self.tmp / "a.csv"
        csvf.write_text("nome;valor\nA;1\nB;2\n", encoding="utf-8")
        ok("L3: B | 2" in m.ler(str(csvf)), "csv com ponto e vírgula")
        d = docx.Document()
        d.add_paragraph("Proposta Raiz")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Plano", "R$ 197"
        dx = self.tmp / "p.docx"
        d.save(dx)
        r = m.ler(str(dx))
        ok("Proposta Raiz" in r and "Plano | R$ 197" in r, f"docx com tabela: {r}")
        w = pypdf.PdfWriter()
        w.add_blank_page(width=200, height=200)
        pdf = self.tmp / "s.pdf"
        with pdf.open("wb") as f:
            w.write(f)
        r = m.ler(str(pdf))
        ok("escaneadas" in r and "pdftoppm" in r, f"pdf sem texto: avisa e diz como ler: {r}")
        txt = self.tmp / "t.txt"
        txt.write_text("\n".join(f"linha {i}" for i in range(1, 11)), encoding="utf-8")
        ok("3: linha 3" in m.ler(str(txt), inicio=3, limite=2) and "4: linha 4" in m.ler(str(txt), inicio=3, limite=2),
           "texto com início e limite")
        ok("não encontrado" in m.ler(str(self.tmp / "x.xlsx")), "arquivo que não existe")
        b = self.tmp / "semext"
        b.write_bytes(b"PK\x03\x04\x00\x00binario")
        ok("binário" in m.ler(str(b)), "binário sem extensão: diz o que fazer")

    def t_ferramenta_tem_manual_e_esquema(self):
        md = (RAIZ / "prompts" / "agent.system.tool.ler.md").read_text(encoding="utf-8")
        ok("valores calculados" in md and '"caminho"' in md, "manual da ferramenta")
        ok("class Ler(Tool)" in (RAIZ / "tools" / "ler.py").read_text(encoding="utf-8"), "ferramenta registrada")

    async def rodar(self, so=""):
        for nome in sorted(n for n in dir(self) if n.startswith("t_") and (not so or so in n)):
            try:
                r = getattr(self, nome)()
                if hasattr(r, "__await__"):
                    await r
                self.resultados.append({"teste": nome, "ok": True})
            except Exception as exc:
                self.resultados.append({"teste": nome, "ok": False, "erro": f"{type(exc).__name__}: {exc}",
                                        "onde": traceback.format_exc().splitlines()[-3:]})
        return {"passou": sum(r["ok"] for r in self.resultados), "total": len(self.resultados),
                "resultados": self.resultados}


if __name__ == "__main__":
    r = asyncio.run(Suite().rodar(sys.argv[1] if len(sys.argv) > 1 else ""))
    for x in r["resultados"]:
        print("ok   " if x["ok"] else "FALHA", x["teste"], x.get("erro", ""), *(x.get("onde") or []) if not x["ok"] else "")
    print("TESTES", r["passou"], r["total"])
