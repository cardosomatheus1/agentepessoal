"""Read a whole file in pages: spreadsheets with their computed values, Word, PDF, CSV and text.

Item 4 of docs/AUDITORIA_MATURIDADE_AGENTE.md. Agent Zero's office reader showed 80 rows x 30 columns x 8 sheets of a
spreadsheet, and with the FORMULAS (data_only=False), not the values: any count or total over a sheet came out wrong,
with no way to read the rest. Hermes (tools/read_extract.py:474-560, tools/file_tools.py:606) reads the cached values,
thousands of rows, skips hidden sheets and pages the output; this does the same, plus a warning on scanned PDF pages
(read_extract.py:291).
"""

import csv
from pathlib import Path

LIMITE_PADRAO = 200      # rows / lines / pages per call
LIMITE_MAX = 2000
CARACTERES_MAX = 20000   # text returned per call; the rest is reached with `inicio`
CELULA_MAX = 200


def _celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).replace("\n", " ").replace("\r", " ").strip()
    return t if len(t) <= CELULA_MAX else t[:CELULA_MAX] + "…"


def _linha(n: int, valores) -> str:
    cel = [_celula(v) for v in valores]
    while cel and not cel[-1]:
        cel.pop()
    return f"L{n}: " + " | ".join(cel) if cel else ""


def _janela(inicio, limite) -> tuple[int, int]:
    try:
        inicio = max(1, int(inicio or 1))
    except (TypeError, ValueError):
        inicio = 1
    try:
        limite = int(limite or LIMITE_PADRAO)
    except (TypeError, ValueError):
        limite = LIMITE_PADRAO
    return inicio, max(1, min(limite, LIMITE_MAX))


def _cortar(linhas: list[str], cabecalho: str, inicio: int, ultimo: int, total: int, unidade: str) -> str:
    """Join, cut at CARACTERES_MAX on a whole line, and say how to read on."""
    saida, tamanho, mostrado = [cabecalho], len(cabecalho), inicio - 1
    for linha in linhas:
        if tamanho + len(linha) + 1 > CARACTERES_MAX and len(saida) > 1:
            break
        saida.append(linha)
        tamanho += len(linha) + 1
        mostrado += 1
    if linhas and mostrado < ultimo:
        ultimo = mostrado
    if ultimo < total:
        saida.append(f"… mostrando {unidade} {inicio}–{ultimo} de {total}. Para continuar: `inicio` = {ultimo + 1}.")
    return "\n".join(saida)


def ler_planilha(caminho: Path, inicio=1, limite=None, aba: str = "") -> str:
    import openpyxl

    inicio, limite = _janela(inicio, limite)
    valores = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        formulas = openpyxl.load_workbook(caminho, read_only=True, data_only=False)
    except Exception:
        formulas = None
    abas = valores.sheetnames
    visiveis = [n for n in abas if valores[n].sheet_state == "visible"]
    if aba:
        alvo = next((n for n in abas if n.lower() == aba.lower()), None)
        if alvo is None:
            return f"Aba {aba!r} não existe. Abas: {', '.join(abas)}."
        escolhidas = [alvo]
    else:
        escolhidas = visiveis
    partes = [f"Planilha {caminho.name}: abas {', '.join(abas)}"
              + (f" (escondidas, fora desta leitura: {', '.join(n for n in abas if n not in visiveis)})"
                 if len(visiveis) < len(abas) and not aba else "")
              + ". Valores calculados (não fórmulas)."]
    for nome in escolhidas:
        ws = valores[nome]
        total = ws.max_row or 0
        ultimo = min(total, inicio + limite - 1)
        linhas, sem_valor = [], 0
        fws = formulas[nome] if formulas else None
        frows = fws.iter_rows(min_row=inicio, max_row=ultimo, values_only=True) if fws is not None else None
        for n, row in enumerate(ws.iter_rows(min_row=inicio, max_row=ultimo, values_only=True), start=inicio):
            frow = next(frows, None) if frows is not None else None
            if frow is not None:
                sem_valor += sum(1 for v, f in zip(row, frow)
                                 if v is None and isinstance(f, str) and f.startswith("="))
            texto = _linha(n, row)
            if texto:
                linhas.append(texto)
        cab = f"\n## Aba {nome!r}: {total} linhas × {ws.max_column or 0} colunas"
        if sem_valor:
            cab += (f"\n⚠️ {sem_valor} célula(s) com fórmula sem valor calculado (o arquivo foi salvo por um programa que "
                    "não calcula). Recalcule antes de somar: `soffice --headless --convert-to xlsx` numa cópia.")
        partes.append(_cortar(linhas, cab, inicio, ultimo, total, "linhas"))
    valores.close()
    if formulas:
        formulas.close()
    return "\n".join(partes)


def ler_xls(caminho: Path, inicio=1, limite=None, aba: str = "") -> str:
    import xlrd

    inicio, limite = _janela(inicio, limite)
    wb = xlrd.open_workbook(str(caminho))
    nomes = wb.sheet_names()
    escolhidas = [n for n in nomes if not aba or n.lower() == aba.lower()]
    if not escolhidas:
        return f"Aba {aba!r} não existe. Abas: {', '.join(nomes)}."
    partes = [f"Planilha {caminho.name}: abas {', '.join(nomes)}."]
    for nome in escolhidas:
        sh = wb.sheet_by_name(nome)
        ultimo = min(sh.nrows, inicio + limite - 1)
        linhas = [t for n in range(inicio, ultimo + 1) if (t := _linha(n, sh.row_values(n - 1)))]
        partes.append(_cortar(linhas, f"\n## Aba {nome!r}: {sh.nrows} linhas × {sh.ncols} colunas", inicio, ultimo,
                              sh.nrows, "linhas"))
    return "\n".join(partes)


def ler_csv(caminho: Path, inicio=1, limite=None) -> str:
    inicio, limite = _janela(inicio, limite)
    texto = caminho.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialeto = csv.excel
    todas = list(csv.reader(texto.splitlines(), dialeto))
    ultimo = min(len(todas), inicio + limite - 1)
    linhas = [t for n in range(inicio, ultimo + 1) if (t := _linha(n, todas[n - 1]))]
    return _cortar(linhas, f"CSV {caminho.name}: {len(todas)} linhas.", inicio, ultimo, len(todas), "linhas")


def ler_docx(caminho: Path, inicio=1, limite=None) -> str:
    import docx

    inicio, limite = _janela(inicio, limite)
    d = docx.Document(str(caminho))
    blocos = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for i, t in enumerate(d.tables, start=1):
        blocos.append(f"[tabela {i}]")
        blocos += [" | ".join(_celula(c.text) for c in r.cells) for r in t.rows]
    ultimo = min(len(blocos), inicio + limite - 1)
    linhas = [f"{n}: {blocos[n - 1]}" for n in range(inicio, ultimo + 1)]
    return _cortar(linhas, f"Documento {caminho.name}: {len(blocos)} parágrafos e linhas de tabela.", inicio, ultimo,
                   len(blocos), "parágrafos")


def ler_pdf(caminho: Path, inicio=1, limite=None) -> str:
    import pypdf

    inicio, limite = _janela(inicio, limite or 20)
    r = pypdf.PdfReader(str(caminho))
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            return "PDF protegido por senha: peça a senha (ou use o cofre) e abra com pypdf decrypt."
    total = len(r.pages)
    ultimo = min(total, inicio + limite - 1)
    linhas, escaneadas = [], []
    for n in range(inicio, ultimo + 1):
        texto = (r.pages[n - 1].extract_text() or "").strip()
        if len(texto) < 20:
            escaneadas.append(n)
        linhas.append(f"--- página {n} ---\n{texto}")
    cab = f"PDF {caminho.name}: {total} páginas."
    if escaneadas:
        cab += (f"\n⚠️ Páginas sem texto (provavelmente escaneadas): {', '.join(map(str, escaneadas))}. Para ler: "
                "`pdftoppm -r 200 -f N -l N -png arquivo.pdf /tmp/p` e `analisar_imagens` (ou `tesseract`) nas imagens.")
    return _cortar(linhas, cab, inicio, ultimo, total, "páginas")


def ler_texto(caminho: Path, inicio=1, limite=None) -> str:
    inicio, limite = _janela(inicio, limite)
    todas = caminho.read_text(encoding="utf-8", errors="replace").splitlines()
    ultimo = min(len(todas), inicio + limite - 1)
    linhas = [f"{n}: {todas[n - 1]}" for n in range(inicio, ultimo + 1)]
    return _cortar(linhas, f"Arquivo {caminho.name}: {len(todas)} linhas.", inicio, ultimo, len(todas), "linhas")


def ler(caminho: str, inicio=1, limite=None, aba: str = "") -> str:
    p = Path(str(caminho or "").strip())
    if not p.is_file():
        return f"Arquivo não encontrado: {p}. Confira o caminho (os anexos ficam em /a0/usr/chats/<conversa>/anexos e os downloads em /a0/usr/downloads)."
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return ler_planilha(p, inicio, limite, aba)
    if ext == ".xls":
        return ler_xls(p, inicio, limite, aba)
    if ext in (".csv", ".tsv"):
        return ler_csv(p, inicio, limite)
    if ext == ".docx":
        return ler_docx(p, inicio, limite)
    if ext == ".pdf":
        return ler_pdf(p, inicio, limite)
    with p.open("rb") as f:
        cabeca = f.read(4096)
    if b"\x00" in cabeca:
        return (f"{p.name} é binário ({ext or 'sem extensão'}). Planilha/Word/PDF sem extensão? Copie com a extensão certa "
                "e leia de novo; imagem: `vision_load` ou `analisar_imagens`.")
    return ler_texto(p, inicio, limite)
